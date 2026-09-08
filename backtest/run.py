"""Backtest runner.

Answers Q1 from METHODOLOGY.md: do exit rules add value on the kind of names
this screener surfaces? Does NOT claim to validate the screen itself — that
needs point-in-time fundamentals.

Usage:
    python -m backtest.run                 # full locked grid
    python -m backtest.run --overlays      # add regime / momentum gates
    python -m backtest.run --quick         # fewer entry dates, for a smoke test
"""

from __future__ import annotations

import argparse
import datetime as dt
import glob
import os
import random
import statistics as stat
import sys

import openpyxl

from . import overlays, prices
from .engine import RuleResult, Trade, simulate, summarize
from .exits import build_grid, buy_and_hold

BENCH = "SPY"
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "output")


def screened_universe() -> list[str]:
    """Every ticker that appeared in any stored screen output."""
    tickers: set[str] = set()
    for path in glob.glob(os.path.join(OUTPUT_DIR, "screener_output_*.xlsx")):
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except (OSError, ValueError):
            continue
        for sheet in ("Nano Cap", "Small Cap", "Breakout"):
            if sheet not in wb.sheetnames:
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            hi = next((i for i, r in enumerate(rows)
                       if r and "Ticker" in [c for c in r if c]), None)
            if hi is None:
                continue
            for r in rows[hi + 1:]:
                if r and r[1] and str(r[0]).isdigit():
                    tickers.add(str(r[1]).strip())
    return sorted(tickers)


def entry_dates(start: dt.date, end: dt.date, step_days: int) -> list[dt.date]:
    out, d = [], start
    while d <= end:
        out.append(d)
        d += dt.timedelta(days=step_days)
    return out


def run_grid(bars: dict[str, prices.Bars], dates: list[dt.date],
             gate=None) -> dict[str, list[Trade]]:
    """Run every rule in the locked grid over every (ticker, date) pair."""
    results: dict[str, list[Trade]] = {}
    grid = build_grid()
    for rule in grid:
        trades: list[Trade] = []
        for tk, b in bars.items():
            for day in dates:
                if gate is not None and not gate(tk, b, day):
                    continue
                t = simulate(b, day, rule)
                if t is not None:
                    trades.append(t)
        results[rule.name] = trades
    return results


def by_year(trades: list[Trade]) -> dict[int, list[float]]:
    out: dict[int, list[float]] = {}
    for t in trades:
        out.setdefault(t.entry_date.year, []).append(t.ret_pct)
    return out


def run_quadrant() -> int:
    """Four-quadrant regime analysis on 20y of index history.

    Uses index proxies rather than the screened universe because the universe
    only has five years of price history — too short for the 10-month SMA to
    produce enough classified months to judge, and short enough that SPY itself
    shows no differentiation across quadrants.
    """
    from . import quadrant as qd

    bars = {t: prices.load(t, rng="20y") for t in ("SPY", "DBC", "IWM", "QQQ")}
    if any(b is None for b in bars.values()):
        print("could not load index history", file=sys.stderr)
        return 1
    qmap = qd.quadrant_by_month(bars["SPY"], bars["DBC"])
    months = sorted(qmap)
    dates = [dt.date(y, m, 1) for (y, m) in months
             if dt.date(y, m, 1) <= dt.date(2025, 9, 1)]
    print(f"classified months: {len(months)}  {months[0]} -> {months[-1]}")
    print(f"entry months with a full 252d forward window: {len(dates)}\n")

    rule = buy_and_hold(252)

    def fwd(tk: str, sel: list[dt.date]) -> list[float]:
        out = []
        for d in sel:
            t = simulate(bars[tk], d, rule)
            if t:
                out.append(t.ret_pct)
        return out

    print("FORWARD 252d RETURN BY QUADRANT AT ENTRY")
    print("="*66)
    print(f"{'quadrant':<16}{'n':>5}{'IWM':>12}{'QQQ':>10}{'SPY':>10}")
    print("-"*66)
    for q in qd.Quadrant:
        sel = [d for d in dates if qd.quadrant_on(qmap, d) == q]
        cells = []
        for tk in ("IWM", "QQQ", "SPY"):
            r = fwd(tk, sel)
            cells.append(f"{stat.mean(r):>+12.1f}" if r else f"{'—':>12}")
        print(f"{q.value:<16}{len(sel):>5}" + "".join(cells))

    # The stagflation stand-down looks adoptable until it is split by era.
    print("\nSTAND DOWN IN STAGFLATION? (small caps)")
    print("-"*66)
    allm = fwd("IWM", dates)
    ex = fwd("IWM", [d for d in dates
                     if qd.quadrant_on(qmap, d) != qd.Quadrant.STAGFLATION])
    print(f"  all months {stat.mean(allm):+.1f}%   ex-stagflation "
          f"{stat.mean(ex):+.1f}%   lift {stat.mean(ex) - stat.mean(allm):+.1f}pp")
    for lab, lo, hi in (("2007-2015", 2007, 2015), ("2016-2026", 2016, 2026)):
        win = [d for d in dates if lo <= d.year <= hi]
        a = fwd("IWM", win)
        e = fwd("IWM", [d for d in win
                        if qd.quadrant_on(qmap, d) != qd.Quadrant.STAGFLATION])
        s = fwd("IWM", [d for d in win
                        if qd.quadrant_on(qmap, d) == qd.Quadrant.STAGFLATION])
        if a and e:
            print(f"  {lab}: lift {stat.mean(e) - stat.mean(a):+.1f}pp   "
                  f"(stagflation n={len(s)}, mean "
                  f"{stat.mean(s) if s else 0:+.1f}%)")
    yrs = sorted({d.year for d in dates
                  if qd.quadrant_on(qmap, d) == qd.Quadrant.STAGFLATION})
    print(f"  stagflation months fall in {yrs} — {len(yrs)} episodes, not "
          f"independent observations")
    print("  -> fails the era-stability standard; not adopted (see FINDINGS.md)")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--quick", action="store_true",
                    help="12-week entry spacing instead of 4-week")
    ap.add_argument("--overlays", action="store_true",
                    help="also test regime and momentum entry gates")
    ap.add_argument("--quadrant", action="store_true",
                    help="four-quadrant macro regime analysis on a long history")
    args = ap.parse_args(argv)

    if args.quadrant:
        return run_quadrant()

    universe = screened_universe()
    print(f"universe: {len(universe)} tickers from stored screens")

    all_bars = prices.load_many(universe + [BENCH])
    bench = all_bars.pop(BENCH, None)
    if bench is None:
        print("could not load benchmark; aborting", file=sys.stderr)
        return 1
    print(f"loaded bars for {len(all_bars)} tickers "
          f"({len(bench)} benchmark bars)\n")

    step = 84 if args.quick else 28
    start = dt.date(2021, 10, 1)
    end = dt.date(2026, 3, 1)   # leaves >=6 months of forward data for every entry
    dates = entry_dates(start, end, step)
    print(f"entries: {len(dates)} dates, every {step} days, {start} to {end}")

    results = run_grid(all_bars, dates)
    baseline_name = buy_and_hold(252).name
    base = summarize(baseline_name, results[baseline_name])
    assert base is not None

    print(f"\n{'='*95}")
    print("Q1: EXIT RULES  (positions are independent, equal-weight, no costs)")
    print("="*95)
    print(RuleResult.header())
    print("-"*95)
    rows = [summarize(k, v) for k, v in results.items()]
    rows = [r for r in rows if r is not None]
    rows.sort(key=lambda r: -r.median)
    for r in rows:
        mark = "  <-- baseline" if r.name == baseline_name else ""
        print(r.row() + mark)

    # ---- locked decision criteria -------------------------------------
    print(f"\n{'='*95}")
    print("DECISION CRITERIA (locked in METHODOLOGY.md before running)")
    print("="*95)
    print(f"baseline = {baseline_name}: median {base.median:+.1f}%, "
          f"ret/dd {base.ret_over_dd:.2f}\n")
    random.seed(42)
    half = set(random.sample(sorted(all_bars), len(all_bars) // 2))

    print(f"{'rule':<18}{'C1 med':>9}{'C2 r/dd':>9}{'C3 yrs':>9}{'C4 split':>10}  verdict")
    print("-"*70)
    verdicts = []
    for r in rows:
        if r.name == baseline_name:
            continue
        trades = results[r.name]
        c1 = r.median > base.median
        c2 = r.ret_over_dd > base.ret_over_dd
        yrs = by_year(trades)
        base_yrs = by_year(results[baseline_name])
        wins = sum(1 for y, v in yrs.items()
                   if y in base_yrs and stat.median(v) > stat.median(base_yrs[y]))
        c3 = wins >= 3
        a = [t.ret_pct for t in trades if t.ticker in half]
        b = [t.ret_pct for t in trades if t.ticker not in half]
        ba = [t.ret_pct for t in results[baseline_name] if t.ticker in half]
        bb = [t.ret_pct for t in results[baseline_name] if t.ticker not in half]
        c4 = bool(a and b and ba and bb and
                  (stat.median(a) - stat.median(ba) > 0) ==
                  (stat.median(b) - stat.median(bb) > 0))
        passed = sum([c1, c2, c3, c4])
        verdict = ("ADOPT" if passed == 4 else
                   "interesting" if passed == 3 else "reject")
        verdicts.append((r.name, passed, verdict))
        print(f"{r.name:<18}{'Y' if c1 else 'n':>9}{'Y' if c2 else 'n':>9}"
              f"{wins:>6}/{len(yrs):<2}{'Y' if c4 else 'n':>10}  {verdict}")

    adopted = [v for v in verdicts if v[2] == "ADOPT"]
    print(f"\n{len(adopted)} rule(s) met all four criteria: "
          f"{', '.join(v[0] for v in adopted) if adopted else 'none'}")

    # ---- disclosed flaw in criterion C1 --------------------------------
    # C1 was locked as "beats buy-and-hold on median". That criterion turns out
    # to be degenerate for threshold exits: if more than half of positions ever
    # touch +X%, a "+X% target" rule reports a median of almost exactly +X% by
    # construction. It measures the threshold, not the rule's quality. This is
    # reported rather than swapped out, because changing the yardstick after
    # seeing results is precisely what the methodology forbids.
    print(f"\n{'='*95}")
    print("DISCLOSED FLAW IN CRITERION C1 — median is degenerate for target rules")
    print("="*95)
    print(f"{'rule':<18}{'median':>9}{'mean':>9}{'vs hold':>10}{'top-decile share':>19}")
    print("-"*66)
    for r in sorted(rows, key=lambda r: -r.mean)[:12]:
        trades = sorted((t.ret_pct for t in results[r.name]), reverse=True)
        top = trades[:max(1, len(trades) // 10)]
        total = sum(t - min(0, min(trades)) for t in trades) or 1.0
        share = 100.0 * sum(t - min(0, min(trades)) for t in top) / total
        mark = "  <-- baseline" if r.name == baseline_name else ""
        print(f"{r.name:<18}{r.median:>+9.1f}{r.mean:>+9.1f}"
              f"{r.mean - base.mean:>+10.1f}{share:>18.0f}%{mark}")
    print("\nRead the mean column, not the median: capping a right-skewed"
          "\ndistribution raises the median while destroying the mean.")

    # ---- overlays -----------------------------------------------------
    if args.overlays:
        print(f"\n{'='*95}")
        print("OVERLAYS (entry gates from the IBKR spec — a layer, not the strategy)")
        print("="*95)

        on_days = [d for d in dates if overlays.regime_active(bench, d)]
        print(f"SPY 50/200 regime ON for {len(on_days)}/{len(dates)} entry dates")
        if len(on_days) in (0, len(dates)):
            print("  -> regime is constant across this window; not testable here.")
        else:
            for label, gate in (
                ("regime ON only",
                 lambda tk, b, d: overlays.regime_active(bench, d)),
                ("12-1 momentum > 0",
                 lambda tk, b, d: (overlays.momentum_12_1(b, d) or -1) > 0),
                ("regime AND momentum",
                 lambda tk, b, d: overlays.regime_active(bench, d)
                 and (overlays.momentum_12_1(b, d) or -1) > 0),
            ):
                trades = [t for tk, b in all_bars.items() for d in dates
                          if gate(tk, b, d)
                          for t in [simulate(b, d, buy_and_hold(252))] if t]
                s = summarize(label, trades)
                if s is None:
                    print(f"  {label:<22} no qualifying entries")
                    continue
                print(f"  {label:<22} n={s.n:<6} mean {s.mean:+6.1f}% "
                      f"(ungated {base.mean:+.1f}%)  median {s.median:+5.1f}% "
                      f"win {s.win_rate:.0f}%  ret/dd {s.ret_over_dd:.2f}")

        # A gate that merely excludes a bad calendar year will look like a good
        # filter in aggregate. The only way to tell the two apart is to compare
        # gated against ungated *within* each year, where both states occur.
        print("\n  Regime ON vs OFF within each year (aggregate lift can be a\n"
              "  composition effect — a gate that just drops one bad year):")
        print(f"  {'year':<6}{'n ON':>7}{'mean ON':>10}{'n OFF':>8}"
              f"{'mean OFF':>11}{'lift':>9}")
        rows_y: list[tuple[int, float]] = []
        graded = [(d, overlays.regime_active(bench, d), t.ret_pct)
                  for tk, b in all_bars.items() for d in dates
                  for t in [simulate(b, d, buy_and_hold(252))] if t]
        for y in sorted({d.year for d, _, _ in graded}):
            on = [r for d, g, r in graded if d.year == y and g]
            off = [r for d, g, r in graded if d.year == y and not g]
            if not on or not off:
                print(f"  {y:<6}{len(on):>7}{'—':>10}{len(off):>8}{'—':>11}"
                      f"{'  regime constant':>9}")
                continue
            lift = stat.mean(on) - stat.mean(off)
            rows_y.append((y, lift))
            print(f"  {y:<6}{len(on):>7}{stat.mean(on):>+10.1f}{len(off):>8}"
                  f"{stat.mean(off):>+11.1f}{lift:>+9.1f}")
        if rows_y:
            wins = sum(1 for _, l in rows_y if l > 0)
            print(f"\n  regime ON beat OFF in {wins}/{len(rows_y)} years with "
                  f"both states present")
            if wins <= len(rows_y) // 2:
                print("  -> the aggregate lift is NOT a regime effect.")

    # ---- is the universe even worth owning? ----------------------------
    print(f"\n{'='*95}")
    print("BENCHMARK: same entry dates, SPY held 252d")
    print("="*95)
    spy_trades = [t for d in dates
                  for t in [simulate(bench, d, buy_and_hold(252))] if t]
    sb = summarize("SPY hold_252d", spy_trades)
    if sb is not None:
        print(f"  SPY      n={sb.n:<5} mean {sb.mean:+6.1f}%  median {sb.median:+6.1f}%  "
              f"win {sb.win_rate:.0f}%  ret/dd {sb.ret_over_dd:.2f}")
        print(f"  screened n={base.n:<5} mean {base.mean:+6.1f}%  median {base.median:+6.1f}%  "
              f"win {base.win_rate:.0f}%  ret/dd {base.ret_over_dd:.2f}")
        print(f"  excess: mean {base.mean - sb.mean:+.1f}pp, "
              f"median {base.median - sb.median:+.1f}pp")
        print("\n  NOTE: the screened universe is selected on 2026 characteristics and\n"
              "  backtested from 2021, so this excess is inflated by construction.\n"
              "  It is an upper bound, not an estimate.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
